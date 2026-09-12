"""File parsing and text chunking.

Supported: PDF, DOCX, XLSX, CSV. Page numbers are preserved for PDFs; DOCX text
is paginated heuristically; spreadsheets are flattened to readable rows so they
are still searchable in the Copilot (numeric analysis lives in ``excel.py``).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

import pymupdf as fitz  # PyMuPDF
from docx import Document as DocxDocument

SUPPORTED_EXT = {"pdf", "docx", "xlsx", "csv"}
_CHARS_PER_PAGE = 1800


class UnsupportedFileType(ValueError):
    pass


class CorruptFile(ValueError):
    pass


@dataclass
class Page:
    page_number: int
    text: str


@dataclass
class ParsedDocument:
    pages: list[Page]

    @property
    def page_count(self) -> int:
        return len(self.pages)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages)


def detect_file_type(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "xls":
        ext = "xlsx"
    if ext not in SUPPORTED_EXT:
        raise UnsupportedFileType(
            f"Unsupported file type '.{ext}'. Allowed: PDF, DOCX, XLSX, CSV."
        )
    return ext


def parse(content: bytes, file_type: str) -> ParsedDocument:
    try:
        if file_type == "pdf":
            return _parse_pdf(content)
        if file_type == "docx":
            return _parse_docx(content)
        if file_type == "csv":
            return _parse_csv(content)
        if file_type == "xlsx":
            return _parse_xlsx(content)
    except (UnsupportedFileType, CorruptFile):
        raise
    except Exception as exc:  # noqa: BLE001
        raise CorruptFile(f"Could not read the file: {exc}") from exc
    raise UnsupportedFileType(file_type)


def _parse_pdf(content: bytes) -> ParsedDocument:
    pages: list[Page] = []
    with fitz.open(stream=content, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append(Page(page_number=i, text=text))
    if not pages:
        raise CorruptFile("No extractable text found in the PDF.")
    return ParsedDocument(pages=pages)


def _paginate(blocks: list[str]) -> list[Page]:
    pages: list[Page] = []
    buf, page_no = "", 1
    for block in blocks:
        if len(buf) + len(block) > _CHARS_PER_PAGE and buf:
            pages.append(Page(page_number=page_no, text=buf.strip()))
            buf, page_no = "", page_no + 1
        buf += block + "\n"
    if buf.strip():
        pages.append(Page(page_number=page_no, text=buf.strip()))
    return pages or [Page(page_number=1, text="")]


def _parse_docx(content: bytes) -> ParsedDocument:
    doc = DocxDocument(io.BytesIO(content))
    blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    if not blocks:
        raise CorruptFile("No text found in the DOCX file.")
    return ParsedDocument(pages=_paginate(blocks))


def _parse_csv(content: bytes) -> ParsedDocument:
    text = content.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)
    if not rows:
        raise CorruptFile("The CSV file is empty.")
    header = rows[0]
    blocks = [", ".join(header)]
    for r in rows[1:]:
        blocks.append("; ".join(f"{h}: {v}" for h, v in zip(header, r)))
    return ParsedDocument(pages=_paginate(blocks))


def _parse_xlsx(content: bytes) -> ParsedDocument:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    blocks: list[str] = []
    for ws in wb.worksheets:
        blocks.append(f"[Sheet: {ws.title}]")
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header:
            blocks.append(", ".join(str(h) for h in header if h is not None))
        for r in rows:
            cells = [str(c) for c in r if c is not None]
            if cells:
                blocks.append("; ".join(cells))
    wb.close()
    if len(blocks) <= 1:
        raise CorruptFile("No data found in the spreadsheet.")
    return ParsedDocument(pages=_paginate(blocks))


# ── chunking ─────────────────────────────────────────────────────────────
@dataclass
class Chunk:
    document_id: str
    user_id: str
    content: str
    page_number: int
    chunk_number: int


def chunk_document(
    parsed: ParsedDocument,
    *,
    document_id: str,
    user_id: str,
    target_chars: int = 900,
    overlap: int = 150,
) -> list[Chunk]:
    chunks: list[Chunk] = []
    n = 0
    for page in parsed.pages:
        text = page.text.strip()
        if not text:
            continue
        start = 0
        while start < len(text):
            end = min(start + target_chars, len(text))
            # try to break on a sentence boundary
            if end < len(text):
                dot = text.rfind(". ", start + target_chars // 2, end)
                if dot != -1:
                    end = dot + 1
            piece = text[start:end].strip()
            if piece:
                n += 1
                chunks.append(
                    Chunk(
                        document_id=document_id,
                        user_id=user_id,
                        content=piece,
                        page_number=page.page_number,
                        chunk_number=n,
                    )
                )
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
    return chunks
