"""Generate realistic sample files (DOCX + XLSX + CSV) for manual upload testing.

Run from the repo root:  python sample-data/generate_samples.py
Requires the backend virtualenv (python-docx, openpyxl).
"""

from __future__ import annotations

import csv
import math
import random
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).parent


def maintenance_report() -> None:
    doc = Document()
    doc.add_heading("Coastal Ridge Wind Farm — Weekly Maintenance Report", 0)
    doc.add_paragraph("Reporting period: Week 37 | Prepared by: Maintenance Team")
    doc.add_heading("1. Fleet summary", level=1)
    doc.add_paragraph(
        "Overall fleet availability was 96.2% this week, down from 98.1%. "
        "Two turbines required unscheduled attention."
    )
    doc.add_heading("2. WTG-3 — gearbox vibration", level=1)
    doc.add_paragraph(
        "Abnormal gearbox vibration was detected on the high-speed shaft "
        "(4.7 mm/s RMS vs a 3.5 mm/s alarm threshold). The turbine was curtailed "
        "to 60% output. The maintenance team must inspect the WTG-3 gearbox with a "
        "borescope before returning it to normal operating conditions. Priority: high."
    )
    doc.add_heading("3. WTG-7 — cooling system alarm", level=1)
    doc.add_paragraph(
        "The cooling system high-temperature alarm triggered twice overnight and "
        "coolant was 15% below minimum. This is a critical issue — the maintenance "
        "team must resolve the cooling system alarm and top up coolant immediately."
    )
    doc.add_heading("4. Other items", level=1)
    doc.add_paragraph(
        "The monthly energy production report is due for management review. "
        "The WTG-1 nacelle anemometer shows a 3% calibration drift and should be "
        "recalibrated at the next service window. Gearbox filter spares are low "
        "and should be replenished."
    )
    doc.save(ROOT / "maintenance" / "Maintenance_Report.docx")


def meeting_minutes() -> None:
    doc = Document()
    doc.add_heading("Project Phoenix — Weekly Steering Meeting", 0)
    doc.add_paragraph("Attendees: Project Team, Operations Lead, Finance")
    doc.add_heading("Schedule", level=1)
    doc.add_paragraph(
        "Phase 2 commissioning is tracking two weeks behind baseline due to a "
        "delayed grid connection permit."
    )
    doc.add_heading("Agreed actions", level=1)
    for line in [
        "Project team to submit the revised grid connection application within one week.",
        "Finance to review the monthly energy report and confirm the revised budget forecast.",
        "Operations to schedule the WTG-3 gearbox inspection alongside crane mobilisation.",
    ]:
        doc.add_paragraph(line, style="List Bullet")
    doc.add_heading("Risks", level=1)
    doc.add_paragraph(
        "If the permit is not granted within three weeks, Phase 2 revenue moves "
        "to the next quarter. Mitigation: escalate to the regulator."
    )
    doc.save(ROOT / "meetings" / "Project_Meeting_Minutes.docx")


def energy_xlsx() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Hourly"
    ws.append(["hour", "demand_mw", "solar_mw", "wind_mw", "price_eur_mwh"])
    random.seed(7)
    for h in range(720):
        demand = 120 + 30 * math.sin(h / 24 * 2 * math.pi) + random.uniform(-8, 8)
        solar = max(0, 45 * math.sin((h % 24) / 24 * math.pi) + random.uniform(-4, 4))
        wind = max(0, 35 + 20 * math.sin(h / 60) + random.uniform(-10, 10))
        price = 40 + (demand - solar - wind) * 0.4 + random.uniform(-5, 5)
        if h == 500:  # inject an anomaly
            demand *= 3
        ws.append([h, round(demand, 2), round(solar, 2), round(wind, 2), round(price, 2)])
    wb.save(ROOT / "spreadsheets" / "Energy_Report.xlsx")


def safety_csv() -> None:
    rows = [
        ["date", "site", "incident_type", "severity", "status"],
        ["2026-08-02", "Coastal Ridge", "Near miss", "Low", "Closed"],
        ["2026-08-11", "Coastal Ridge", "Equipment fault", "Medium", "Closed"],
        ["2026-08-19", "Coastal Ridge", "Cooling system alarm", "High", "Open"],
        ["2026-08-27", "Coastal Ridge", "Working at height deviation", "Medium", "Open"],
        ["2026-09-03", "Coastal Ridge", "Gearbox vibration alarm", "High", "Open"],
    ]
    with open(ROOT / "reports" / "Safety_Incidents.csv", "w", newline="") as f:
        csv.writer(f).writerows(rows)


if __name__ == "__main__":
    for sub in ("maintenance", "meetings", "reports", "spreadsheets"):
        (ROOT / sub).mkdir(exist_ok=True)
    maintenance_report()
    meeting_minutes()
    energy_xlsx()
    safety_csv()
    print("Sample files written to sample-data/")
